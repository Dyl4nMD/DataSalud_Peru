MONGO_URI = "mongodb://localhost:27017/"   
DB_NAME = "DIRESA_PERU"
COLLECTION_NAME = "Recien_Nacidos"

RUTA_PRESTACIONES = "1 Prestaciones.xlsx"
RUTA_DIAGNOSTICOS = "2 Diagnósticos.xlsx"
RUTA_MEDICAMENTOS = "3 Medicamentos.xlsx"
RUTA_CATALOGOS = "4 Catálogos.xlsx"

BATCH_SIZE = 5000  


import sys
import time
import openpyxl
from pymongo import MongoClient, UpdateOne


def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def cargar_catalogo_diagnostico(ruta):
    log("Cargando catálogo de diagnósticos...")
    wb = openpyxl.load_workbook(ruta, read_only=True)
    ws = wb["Diagnóstico"]
    catalogo = {}
    for coddia, nombre in ws.iter_rows(min_row=2, values_only=True):
        if coddia is not None:
            catalogo[str(coddia).strip()] = nombre
    wb.close()
    log(f"  {len(catalogo):,} diagnósticos en catálogo.")
    return catalogo


def cargar_catalogo_medicamento(ruta):
    log("Cargando catálogo de medicamentos...")
    wb = openpyxl.load_workbook(ruta, read_only=True)
    ws = wb["Medicamentos"]
    catalogo = {}
    for cod, nombre, presentacion, concentracion, forma in ws.iter_rows(min_row=2, values_only=True):
        if cod is not None:
            catalogo[str(cod).strip()] = {
                "nombre_medicamento": nombre,
                "presentacion": presentacion,
                "concentracion": concentracion,
                "forma_farmaceutica": forma,
            }
    wb.close()
    log(f"  {len(catalogo):,} medicamentos en catálogo.")
    return catalogo


def normalizar_cod_medicamento(valor):
    if valor is None:
        return None
    if isinstance(valor, float):
        valor = int(valor)
    return str(valor).zfill(5)


def cargar_medicamentos_por_episodio_diagnostico(ruta, cat_medicamento):
    log("Cargando y agrupando consumo de medicamentos (archivo más pesado, toma más tiempo)...")
    wb = openpyxl.load_workbook(ruta, read_only=True)
    ws = wb["3 Medicamentos"]
    agrupado = {}
    total = 0
    no_encontrados = 0
    for fecha_corte, id_registro_rel, coddia, cod_medicamento, cantidad, valor_neto in ws.iter_rows(
        min_row=2, values_only=True
    ):
        total += 1
        clave = (id_registro_rel, coddia)
        cod_str = normalizar_cod_medicamento(cod_medicamento)
        info_catalogo = cat_medicamento.get(cod_str)
        if info_catalogo is None:
            no_encontrados += 1
            info_catalogo = {
                "nombre_medicamento": None,
                "presentacion": None,
                "concentracion": None,
                "forma_farmaceutica": None,
            }

        medicamento = {
            "cod_medicamento": cod_str,
            "cantidad_entregada": cantidad,
            "valor_neto": valor_neto,
            **info_catalogo,
        }
        agrupado.setdefault(clave, []).append(medicamento)

        if total % 200000 == 0:
            log(f"  ...{total:,} filas de medicamentos procesadas")
    wb.close()
    log(f"  {total:,} filas de medicamentos procesadas en total.")
    if no_encontrados:
        log(f"  AVISO: {no_encontrados:,} códigos de medicamento no se encontraron en el catálogo.")
    return agrupado


def cargar_diagnosticos_por_episodio(ruta, cat_diagnostico, medicamentos_por_clave):
    log("Cargando y agrupando diagnósticos...")
    wb = openpyxl.load_workbook(ruta, read_only=True)
    ws = wb["2 Diagnósticos"]
    agrupado = {}
    total = 0
    for fecha_corte, id_registro_rel, coddia, tipo_diagnostico in ws.iter_rows(min_row=2, values_only=True):
        total += 1
        coddia_norm = str(coddia).strip() if coddia is not None else None
        diagnostico = {
            "coddia": coddia_norm,
            "nombre_diagnostico": cat_diagnostico.get(coddia_norm),
            "tipo_diagnostico": tipo_diagnostico,
            "medicamentos": medicamentos_por_clave.get((id_registro_rel, coddia), []),
        }
        agrupado.setdefault(id_registro_rel, []).append(diagnostico)
    wb.close()
    log(f"  {total:,} filas de diagnósticos procesadas.")
    return agrupado


def construir_documento(fila_prestacion, diagnosticos_por_episodio):
    (
        fecha_corte,
        id_registro_rel,
        codigo_anonimizado,
        departamento,
        provincia,
        distrito,
        ubigeo,
        nivel_eess,
        codigo_serv_prestacional,
        servicio_prestacional,
        dias_hosp,
        edad_dias,
        fecha_atencion,
        sexo,
        peso,
        ate_capita,
    ) = fila_prestacion

    documento = {
        "_id": str(id_registro_rel),
        "fecha_corte": fecha_corte,
        "afiliado": {
            "codigo_anonimizado": codigo_anonimizado,
            "sexo": sexo,
            "edad_dias": edad_dias,
            "peso_kg": peso,
        },
        "atencion": {
            "fecha_atencion": fecha_atencion,
            "dias_hospitalizacion": dias_hosp,
            "es_capita": (ate_capita == "S"),
            "establecimiento": {
                "ubigeo": str(ubigeo).zfill(6) if ubigeo is not None else None,
                "departamento": departamento,
                "provincia": provincia,
                "distrito": distrito,
                "nivel_eess": str(nivel_eess) if nivel_eess is not None else None,
            },
            "servicio_prestacional": {
                "codigo": str(codigo_serv_prestacional).zfill(3) if codigo_serv_prestacional is not None else None,
                "descripcion": servicio_prestacional,
            },
        },
        "diagnosticos": diagnosticos_por_episodio.get(id_registro_rel, []),
    }
    return documento


def main():
    cat_diagnostico = cargar_catalogo_diagnostico(RUTA_CATALOGOS)
    cat_medicamento = cargar_catalogo_medicamento(RUTA_CATALOGOS)

    medicamentos_por_clave = cargar_medicamentos_por_episodio_diagnostico(RUTA_MEDICAMENTOS, cat_medicamento)
    diagnosticos_por_episodio = cargar_diagnosticos_por_episodio(
        RUTA_DIAGNOSTICOS, cat_diagnostico, medicamentos_por_clave
    )

    del medicamentos_por_clave

    log("Conectando a MongoDB...")
    client = MongoClient(MONGO_URI)
    coleccion = client[DB_NAME][COLLECTION_NAME]

    log("Creando índice sobre afiliado.codigo_anonimizado (no único, un bebé puede tener varios episodios)...")
    coleccion.create_index("afiliado.codigo_anonimizado")

    log("Leyendo Prestaciones y armando/insertando documentos por lotes...")
    wb = openpyxl.load_workbook(RUTA_PRESTACIONES, read_only=True)
    ws = wb["1 Prestaciones"]

    lote = []
    total_insertados = 0
    total_actualizados = 0

    for fila in ws.iter_rows(min_row=2, values_only=True):
        doc = construir_documento(fila, diagnosticos_por_episodio)
        # upsert por _id: si vuelves a correr el script, no duplica documentos
        lote.append(UpdateOne({"_id": doc["_id"]}, {"$set": doc}, upsert=True))

        if len(lote) >= BATCH_SIZE:
            resultado = coleccion.bulk_write(lote, ordered=False)
            total_insertados += resultado.upserted_count
            total_actualizados += resultado.modified_count
            log(f"  ...{total_insertados + total_actualizados:,} documentos procesados")
            lote = []

    if lote:
        resultado = coleccion.bulk_write(lote, ordered=False)
        total_insertados += resultado.upserted_count
        total_actualizados += resultado.modified_count

    wb.close()
    log(f"Listo. Insertados nuevos: {total_insertados:,} | Actualizados (ya existían): {total_actualizados:,}")
    log(f"Total de documentos en la colección: {coleccion.count_documents({}):,}")


if __name__ == "__main__":
    main()
